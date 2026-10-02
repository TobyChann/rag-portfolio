# backend/app/sync/resume.py
"""Resume upload → clean text the chatbot reads on every question. The original
file is kept in Postgres for visitors to download.

The resume is deliberately NOT chunked into the vector DB: it's ~1-2k tokens,
so it goes into the model's context whole. Vector search over it kept missing
entries (a chunk holding just the name outscored the jobs for "work
experience"), and giving the model the full text costs a fraction of a cent."""
import io
import re

from sqlalchemy import select

from app.db.database import AsyncSessionLocal
from app.db.models import Resume

PDF = "application/pdf"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
ALLOWED_TYPES = {".pdf": PDF, ".docx": DOCX}
MAX_BYTES = 5 * 1024 * 1024

# Common resume section titles, promoted to Markdown headings for structure.
_SECTIONS = {
    "education", "experience", "professional experience", "work experience",
    "projects", "personal/ academic projects", "personal / academic projects",
    "personal projects", "academic projects", "skills", "skills & interests",
    "skills and interests", "technical skills", "achievements", "awards",
    "leadership", "activities", "certifications", "co-curricular activities",
}

# Contact details stay in the downloadable file but out of the chatbot's
# context, so it never recites a phone number or email.
_CONTACT = [
    re.compile(r"[\w.+-]+@[\w-]+\.[\w.]+"),                 # email
    re.compile(r"\+?\d[\d\s-]{6,}\d"),                      # phone number
    re.compile(r"https?://\S+|(?:www\.)?linkedin\.com/\S*", re.I),
]


def extract_text(data: bytes, content_type: str) -> str:
    if content_type == PDF:
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    import docx
    doc = docx.Document(io.BytesIO(data))
    lines = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            lines.append(" | ".join(dict.fromkeys(c.text.strip() for c in row.cells)))
    return "\n".join(lines)


def to_markdown(text: str) -> str:
    """Strip contact details and turn section titles into ## headings."""
    out = []
    for line in text.splitlines():
        for pattern in _CONTACT:
            line = pattern.sub("", line)
        line = line.strip(" |\t")
        if not re.sub(r"[\s|]", "", line):   # only separators left once contacts are gone
            continue
        if line.lower().rstrip(":") in _SECTIONS:
            out.append(f"\n## {line.title()}\n")
        else:
            out.append(line)
    return "\n".join(out).strip()


# In-process cache of the current resume text; None = not loaded yet.
# Upload/delete call invalidate_resume_cache().
_cached_text: str | None = None


async def get_resume_text() -> str:
    """Current resume text ("" if none), read from Postgres once and cached."""
    global _cached_text
    if _cached_text is None:
        async with AsyncSessionLocal() as db:
            _cached_text = (await db.execute(select(Resume.text))).scalars().first() or ""
    return _cached_text


def invalidate_resume_cache():
    global _cached_text
    _cached_text = None
